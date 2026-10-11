from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
source = BASE / 'INFORME_PRACTICAS_REDACCION_V02.md'
target = BASE / 'Informe_Practicas_Brayni_Chavez_V02.docx'
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.85)
section.left_margin = section.right_margin = Inches(0.9)
section.footer_distance = Inches(0.35)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Number', 'List Bullet']:
    style = document.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.15
document.styles['Title'].font.size = Pt(21)
document.styles['Subtitle'].font.size = Pt(15)
document.styles['Subtitle'].font.italic = False
for style in document.styles:
    if style._element.pPr is not None:
        for border in list(style._element.pPr.findall(qn('w:pBdr'))):
            style._element.pPr.remove(border)
document.styles['Heading 1'].font.size = Pt(15)
document.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1', 'Heading 2']:
    document.styles[name].font.bold = True
    document.styles[name].paragraph_format.keep_with_next = True
    document.styles[name].paragraph_format.space_before = Pt(15)
    document.styles[name].paragraph_format.space_after = Pt(7)
document.core_properties.title = 'Informe de prácticas preprofesionales'
document.core_properties.author = 'Brayni Chavez Caruajulca'
document.core_properties.subject = 'Sistema de acompañamiento tutorial FISME UNTRM'

def cover(text, style='Normal', before=0, after=10):
    p = document.add_paragraph(text, style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    borders = OxmlElement('w:pBdr')
    for side in ['top', 'left', 'bottom', 'right', 'between', 'bar']:
        border = OxmlElement('w:' + side)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    p._p.get_or_add_pPr().append(borders)
    return p

cover('UNIVERSIDAD NACIONAL TORIBIO RODRÍGUEZ\nDE MENDOZA DE AMAZONAS', before=15)
cover('FACULTAD DE INGENIERÍA DE SISTEMAS\nY MECÁNICA ELÉCTRICA')
cover('ESCUELA PROFESIONAL DE INGENIERÍA DE SISTEMAS')
cover('Informe de prácticas preprofesionales', 'Title', before=30, after=20)
cover('Sistema de Información para la Sistematización de los Procesos de Acompañamiento de Tutoría Universitaria de la Facultad de Ingeniería de Sistemas y Mecánica Eléctrica', 'Subtitle', after=30)
cover('ESTUDIANTE\nBrayni Chavez Caruajulca', after=20)
cover('ASESOR\nMg. Eder Nicanor Figueroa Piscoya', after=25)
cover('Bagua, Amazonas\n2026', after=20)
cover('Versión de trabajo del informe final\nCorte al 11 de octubre de 2026')
document.add_page_break()

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def table(rows):
    t = document.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths_by_header = {
        'Fase': [2.6, 3.4, 0.7],
        'Versión': [0.8, 1.0, 4.9],
        'Verificación': [1.9, 2.9, 1.9],
        'Objetivo': [1.2, 2.8, 2.7],
    }
    widths = widths_by_header.get(rows[0][0], [1.1, 2.5, 3.1])
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:' + edge)
        for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): el.set(qn('w:'+k),v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for r_index, values in enumerate(rows):
        cells = t.rows[0].cells if r_index == 0 else t.add_row().cells
        props = cells[0]._tc.getparent().get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if r_index == 0: props.append(OxmlElement('w:tblHeader'))
        for i,(cell,value) in enumerate(zip(cells,values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '17365D' if r_index == 0 else ('F2F5F8' if r_index%2 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shading)
            margin = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); margin.append(el)
            cell._tc.get_or_add_tcPr().append(margin)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
            if rows[0][0]=='Fase' and i>0: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size=Pt(10)
                run.bold=r_index==0
                run.font.color.rgb=RGBColor(255,255,255) if r_index==0 else RGBColor(0,0,0)
    document.add_paragraph().paragraph_format.space_after=Pt(2)

lines = source.read_text(encoding='utf-8').splitlines()
start = next(i for i,line in enumerate(lines) if line.startswith('## '))
i=start
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
            i+=1
        table(rows); continue
    if line.startswith('### '): document.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('## '): document.add_paragraph(line[3:],'Heading 1')
    elif re.match(r'^\d+\. ',line):
        p=document.add_paragraph(line)
        p.paragraph_format.left_indent=Inches(0.2)
        p.paragraph_format.first_line_indent=Inches(-0.2)
    else:
        p=document.add_paragraph(line)
        p.paragraph_format.widow_control=True
        if line.startswith('['):
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_together=True
            for run in p.runs: run.font.size=Pt(10.5)
    i+=1
document.save(target)
print(str(target))
print('Párrafos:',len(document.paragraphs),'Tablas:',len(document.tables))
